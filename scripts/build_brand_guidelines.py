from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "artifacts" / "brand-guidelines"
OUT_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT_DIR / "Creatorshop Brand Guidelines Book.docx"
BOARD_PATH = OUT_DIR / "brand-board.png"
SYMBOL_PATH = OUT_DIR / "creatorshop-symbol.png"
METAL_PATH = OUT_DIR / "creatorshop-metal-symbol.png"


COLORS = {
    "ink": "171717",
    "black": "0F0F0F",
    "lime": "A3FF38",
    "lime_border": "82F200",
    "surface": "FAFAFA",
    "muted": "737373",
    "line": "E5E5E5",
    "danger": "DC2626",
}


def hex_to_rgb(value: str) -> tuple[int, int, int]:
    value = value.strip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "8") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")


def set_para_spacing(paragraph, before=0, after=8, line=1.08) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_run_font(run, name="Arial", size=None, bold=None, color=None) -> None:
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*hex_to_rgb(color))


def add_text(paragraph, text, size=10.5, bold=False, color=COLORS["ink"], font="Arial") -> None:
    run = paragraph.add_run(text)
    set_run_font(run, name=font, size=size, bold=bold, color=color)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)

    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(32)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(20)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 3"].font.size = Pt(11)
    doc.styles["Heading 3"].font.bold = True


def paragraph(doc, text="", size=10.5, bold=False, color=COLORS["ink"], align=None, before=0, after=8):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    set_para_spacing(p, before=before, after=after)
    if text:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_para_spacing(p, before=12 if level == 1 else 8, after=8)
    for run in p.runs:
        set_run_font(run, size=20 if level == 1 else 14, bold=True, color="000000")
    return p


def make_images() -> None:
    symbol_src = ROOT / "public" / "Creatorshop Brand Symbol.webp"
    metal_src = ROOT / "public" / "Metal Logo.webp"
    for src, dst in ((symbol_src, SYMBOL_PATH), (metal_src, METAL_PATH)):
        img = Image.open(src).convert("RGBA")
        img.thumbnail((1000, 1000))
        canvas = Image.new("RGBA", (1100, 1100), (255, 255, 255, 0))
        canvas.alpha_composite(img, ((1100 - img.width) // 2, (1100 - img.height) // 2))
        canvas.save(dst)

    board = Image.new("RGB", (1800, 1040), hex_to_rgb("FFFFFF"))
    draw = ImageDraw.Draw(board)
    try:
        font_big = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 90)
        font_mid = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 38)
        font_small = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 28)
        font_bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 34)
    except OSError:
        font_big = font_mid = font_small = font_bold = ImageFont.load_default()

    draw.rectangle((0, 0, 1800, 1040), fill=hex_to_rgb("FFFFFF"))
    draw.rectangle((90, 88, 1710, 952), outline=hex_to_rgb("E5E5E5"), width=4)
    draw.rectangle((90, 88, 1710, 378), fill=hex_to_rgb("A3FF38"))
    draw.text((160, 148), "Creatorshop", fill=hex_to_rgb("171717"), font=font_big)
    draw.text((164, 258), "Trade posts for software access.", fill=hex_to_rgb("171717"), font=font_mid)
    draw.text((164, 310), "A self serve barter marketplace for creator made content.", fill=hex_to_rgb("171717"), font=font_small)

    swatches = [
        ("Lime", "A3FF38"),
        ("Lime border", "82F200"),
        ("Ink", "171717"),
        ("Paper", "FFFFFF"),
        ("Soft surface", "FAFAFA"),
        ("Line", "E5E5E5"),
    ]
    x = 160
    for name, value in swatches:
        y = 475
        draw.rounded_rectangle((x, y, x + 205, y + 150), radius=18, fill=hex_to_rgb(value), outline=hex_to_rgb("D9D9D9"), width=2)
        text_color = hex_to_rgb("FFFFFF") if value in ("171717",) else hex_to_rgb("171717")
        draw.text((x + 20, y + 34), name, fill=text_color, font=font_bold)
        draw.text((x + 20, y + 88), "#" + value, fill=text_color, font=font_small)
        x += 245

    draw.rounded_rectangle((160, 720, 710, 840), radius=30, fill=hex_to_rgb("171717"))
    draw.text((210, 754), "Primary action", fill=hex_to_rgb("FFFFFF"), font=font_bold)
    draw.rounded_rectangle((760, 720, 1310, 840), radius=30, fill=hex_to_rgb("A3FF38"), outline=hex_to_rgb("82F200"), width=4)
    draw.text((810, 754), "Accent surface", fill=hex_to_rgb("171717"), font=font_bold)
    draw.text((160, 895), "Use lime as energy, not as paragraph text.", fill=hex_to_rgb("737373"), font=font_mid)
    board.save(BOARD_PATH)


def add_logo_lockup(doc) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell in table.row_cells(0):
        set_cell_border(cell, "FFFFFF", "0")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_width(table.cell(0, 0), 1.45)
    set_cell_width(table.cell(0, 1), 5.8)
    p = table.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SYMBOL_PATH), width=Inches(1.0))
    p = table.cell(0, 1).paragraphs[0]
    set_para_spacing(p, after=0)
    add_text(p, "Creatorshop", size=32, bold=True, color=COLORS["ink"])
    p = table.cell(0, 1).add_paragraph()
    set_para_spacing(p, after=0)
    add_text(p, "Brand guidelines book", size=12, color=COLORS["muted"])


def add_kv_table(doc, rows) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.35, 1.25, 4.45]
    headers = ["Element", "Value", "Guideline"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, COLORS["ink"])
        set_cell_border(cell)
        p = cell.paragraphs[0]
        add_text(p, h, size=9.5, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_para_spacing(p, after=0, line=1.15)
            add_text(p, value, size=9.2, bold=(i == 0), color=COLORS["ink"] if i != 1 else COLORS["muted"])
    paragraph(doc, "", after=4)


def add_color_table(doc) -> None:
    rows = [
        ("Primary lime", "#A3FF38", "Use for accent surfaces, badges, activation states, and moments of energy."),
        ("Lime border", "#82F200", "Use as the structural edge around lime buttons and panels."),
        ("Ink", "#171717", "Use for main text, primary controls, and high contrast system surfaces."),
        ("Soft surface", "#FAFAFA", "Use for FAQ blocks, cards, empty states, and secondary UI panels."),
        ("Line", "#E5E5E5", "Use for borders and quiet dividers."),
        ("Muted text", "#737373", "Use for supporting explanations and metadata."),
        ("Danger", "#DC2626", "Reserve for destructive or error states only."),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.0, 1.55, 1.25, 3.25]
    headers = ["Swatch", "Name", "Hex", "Usage"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, COLORS["ink"])
        set_cell_border(cell)
        add_text(cell.paragraphs[0], h, size=9.2, bold=True, color="FFFFFF")
    for name, value, use in rows:
        cells = table.add_row().cells
        for i, width in enumerate(widths):
            set_cell_width(cells[i], width)
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], value.strip("#"))
        cells[0].paragraphs[0].add_run(" ")
        add_text(cells[1].paragraphs[0], name, size=9.2, bold=True)
        add_text(cells[2].paragraphs[0], value, size=9.2, color=COLORS["muted"])
        add_text(cells[3].paragraphs[0], use, size=9.2)


def add_dos_table(doc, title, rows) -> None:
    heading(doc, title, 2)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(("Do", "Do not")):
        cell = table.cell(0, i)
        set_cell_width(cell, 3.48)
        set_cell_shading(cell, COLORS["ink"] if i == 0 else "F2F2F2")
        set_cell_border(cell)
        add_text(cell.paragraphs[0], h, size=9.5, bold=True, color="FFFFFF" if i == 0 else COLORS["ink"])
    for good, bad in rows:
        cells = table.add_row().cells
        for i, value in enumerate((good, bad)):
            set_cell_width(cells[i], 3.48)
            set_cell_border(cells[i])
            set_cell_shading(cells[i], "FFFFFF")
            p = cells[i].paragraphs[0]
            set_para_spacing(p, after=0, line=1.15)
            add_text(p, value, size=9.2)


def add_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p, after=0)
    add_text(p, "Creatorshop brand guidelines", size=8.5, color=COLORS["muted"])


def build() -> None:
    make_images()
    doc = Document()
    style_doc(doc)
    add_footer(doc.sections[0])

    add_logo_lockup(doc)
    paragraph(doc, "A practical guide for making Creatorshop feel consistent across product, outreach, content, and future brand surfaces.", size=12.5, color=COLORS["ink"], after=18)
    p = paragraph(doc, "", after=14)
    p.add_run().add_picture(str(BOARD_PATH), width=Inches(6.9))
    paragraph(doc, "This book is based on the live codebase, product brief, design system entries, and committed brand assets in the repository.", size=9.5, color=COLORS["muted"], after=4)

    doc.add_page_break()
    heading(doc, "Brand Foundation", 1)
    paragraph(doc, "Creatorshop is a self serve barter marketplace where creators trade posts for software access and software brands trade product access for creator made content. The brand should feel direct, useful, and a little defiant: the creator has value before anyone hands them a media kit or gatekeeps them by follower count.")
    add_kv_table(doc, [
        ("Name", "Creatorshop", "Write as one word with an initial capital. Avoid Creator Shop, CreatorShop, and The Creatorshop."),
        ("Audience", "Creators first", "Prioritize creator clarity while still making brand side value obvious."),
        ("Category", "Software drops", "At launch, describe the marketplace as scheduled software drops rather than a fully stocked always on catalog."),
        ("Mechanism", "Barter", "Creators pay with posts and brands pay with product access. Avoid language that implies cash payouts to creators."),
        ("Core unit", "Shop", "A Shop is the barter workflow: apply, approve, access reveal, deliver, complete."),
    ])
    paragraph(doc, "Positioning: Product Hunt, but creators shop for software by posts.", size=13, bold=True, after=4)
    paragraph(doc, "Use this phrase internally to keep the product sharp. In public copy, unpack it in plain language so people understand the action: browse a drop, apply to a software listing, and pay with content if approved.")

    heading(doc, "Logo System", 1)
    paragraph(doc, "The Creatorshop mark is a node symbol inside a rounded square. It can be used as the full wordmark from Logo.svg, as a standalone symbol, or as a 3D chrome object in high energy product moments.")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (img, label, note) in enumerate(((SYMBOL_PATH, "Primary symbol", "Black node on lime. Use on white and pale neutral surfaces."), (METAL_PATH, "Metal symbol", "Use sparingly for premium, cinematic, or launch visuals."))):
        cell = table.cell(0, idx)
        set_cell_width(cell, 3.45)
        set_cell_border(cell)
        set_cell_shading(cell, "FFFFFF")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Inches(1.7))
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, size=10, bold=True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, note, size=8.7, color=COLORS["muted"])
    doc.add_page_break()
    add_dos_table(doc, "Logo Rules", [
        ("Keep generous clear space around the symbol and wordmark.", "Stretch, rotate, crop, or place the mark in a busy image area."),
        ("Use lime on white, black on lime, or the supplied full wordmark.", "Invent new symbol colors or recolor the wordmark casually."),
        ("Use the 3D symbol for scroll moments, launch material, and hero scenes.", "Use the 3D symbol as a tiny favicon replacement or inside dense forms."),
    ])

    heading(doc, "Color", 1)
    paragraph(doc, "Creatorshop is mostly black, white, and quiet neutral UI. Lime is the signature burst: it signals energy, action, approval, or the marketplace moment. The codebase includes a global guardrail that prevents green copy from returning, so lime should not be used as normal text.")
    add_color_table(doc)
    paragraph(doc, "Practical ratio: default to about 70 percent white or soft surface, 20 percent ink and line work, and 10 percent lime. Landing moments may use more lime, but transactional pages should stay calmer.")

    doc.add_page_break()
    heading(doc, "Typography", 1)
    paragraph(doc, "Satoshi is the brand sans and should be the default for app UI and marketing prose. Plus Jakarta Sans is the closest working substitute when Satoshi is unavailable in production tools. The app also uses expressive supporting fonts for specific product moments.")
    add_kv_table(doc, [
        ("Primary", "Satoshi", "Use for body text, navigation, buttons, cards, and product UI."),
        ("Fallback", "Plus Jakarta Sans", "Use in tools that do not have Satoshi available."),
        ("Mono", "Space Mono", "Use for small technical copy, local time, receipt language, and system details."),
        ("Pixel", "Geist Pixel", "Use sparingly for hero statements, glitchy product moments, and launch energy."),
        ("Handwritten", "Caveat or Permanent Marker", "Use only where the product intentionally feels annotated, creator made, or rough drafted."),
    ])
    paragraph(doc, "Type behavior should be compact and readable. Product screens prefer tight hierarchy, generous line height, and no negative letter spacing in compact controls. Hero scale belongs to true hero moments, not dashboards or dense transactional surfaces.")

    heading(doc, "Voice and Messaging", 1)
    paragraph(doc, "The voice is plainspoken, creator first, and self serve. It should make the barter mechanism feel obvious and legitimate without sounding corporate or inflated.")
    add_kv_table(doc, [
        ("Brand idea", "Your content is the new cash", "Use as a campaign line or hero expression, not as a legal explanation."),
        ("Plain promise", "Trade posts for software access", "Use when readers need to understand the product fast."),
        ("Creator CTA", "Browse software drops", "Use drop language before launch or when inventory is scheduled."),
        ("Workflow", "Apply, approve, access reveal, deliver, complete", "Use this sequence when explaining the Shop lifecycle."),
        ("Brand side", "Creator made content for product access", "Keep the exchange clear. Brands provide access; creators provide agreed content."),
    ])
    doc.add_page_break()
    add_dos_table(doc, "Copy Rules", [
        ("Say creators pay with posts, content, or deliverables.", "Say creators get paid in cash unless a real cash product exists."),
        ("Say no minimum follower count when explaining access.", "Promise approval, traction, audience growth, or guaranteed brand deals."),
        ("Say software drops before marketplace supply is live.", "Pretend there are live listings, testimonials, or traction numbers."),
        ("Use direct verbs like browse, apply, unlock, deliver.", "Use vague verbs like empower, revolutionize, or transform without proof."),
    ])

    heading(doc, "Product UI", 1)
    paragraph(doc, "The interface should feel like a useful marketplace, not a decorative landing page. Product surfaces are neutral, compact, and structured by borders. Motion and playful effects appear in hero or showcase moments, while checkout, settings, messages, and Shop status views stay restrained.")
    add_kv_table(doc, [
        ("Cards", "White or #FAFAFA", "Use 12 to 20 px radii for product cards and transactional surfaces."),
        ("Buttons", "Ink or lime", "Use ink for primary high contrast actions; use lime when action energy matters."),
        ("Borders", "#E5E5E5", "Use borders to create structure before adding shadows."),
        ("Elevation", "Minimal", "Avoid dashboard style heavy shadows. The codebase actively removes some older shadows on creator pages."),
        ("Status", "Badges and dots", "Use lime for approved, active, and ready states; keep negative states clear and restrained."),
    ])

    doc.add_page_break()
    heading(doc, "Imagery and Motion", 1)
    paragraph(doc, "Creatorshop can use visual assets with a strong product signal: the symbol, 3D logo models, creator portraits, software logos, receipt or access metaphors, and interface captures. Pure atmosphere should not replace real product meaning.")
    add_kv_table(doc, [
        ("Hero effects", "Pixel trail and text scramble", "Use for first impression energy and creator culture references."),
        ("3D models", "Symbol, floppy disk, pitch bag, deliver symbol", "Use for story beats about applying, pitching, and delivering."),
        ("ASCII fire", "Lime background", "Use for high energy launch moments, especially the subscription ditching message."),
        ("Creator images", "Square portraits", "Use for profile and directory contexts, not fake testimonials."),
        ("Software logos", "Real product marks in mocks", "Use only when representing mock examples; avoid implying partnerships."),
    ])
    paragraph(doc, "Motion should support comprehension. Press states can scale slightly, hover states can brighten, and hero elements can scramble or react to cursor movement. Avoid adding motion to dense lists where it slows repeated work.")

    heading(doc, "Naming and Terminology", 1)
    paragraph(doc, "Use a small vocabulary consistently so the product feels understandable from the first session.")
    add_kv_table(doc, [
        ("Shop", "One barter transaction", "Capitalize when referring to the product workflow or a user specific deal."),
        ("Drop", "Scheduled listing release", "Use before full supply exists and for launch campaigns."),
        ("Access reveal", "Approved subscription access", "Use after brand approval when the creator unlocks the product details."),
        ("Deliverable", "The promised content", "Use in contracts, product UI, and review states."),
        ("Proof", "Evidence of posting", "Use after a creator publishes or submits completion evidence."),
    ])

    doc.add_page_break()
    heading(doc, "Governance Checklist", 1)
    paragraph(doc, "Use this checklist before shipping public pages, outreach material, design system changes, or new product screens.")
    checklist = [
        "The page explains barter without implying cash payouts.",
        "The copy does not fabricate live listings, partner logos, testimonials, or traction.",
        "Lime is used as surface or action energy, not paragraph text.",
        "Satoshi or the approved fallback drives the main UI.",
        "A Shop means apply, approve, access reveal, deliver, complete.",
        "The creator side feels free and self serve; the brand side remains clear about the $50 per month tool.",
        "Product surfaces prioritize borders, spacing, and compact hierarchy over decorative shadows.",
        "Any playful motion is confined to moments that benefit from it.",
    ]
    for item in checklist:
        p = doc.add_paragraph(style=None)
        set_para_spacing(p, after=5)
        add_text(p, "□ ", size=11, color=COLORS["ink"])
        add_text(p, item, size=10.2, color=COLORS["ink"])

    paragraph(doc, "", after=2)
    paragraph(doc, "Source basis: PRODUCT.md, src/app/globals.css, src/app/layout.tsx, component registry entries, app copy, and assets under public.", size=8.5, color=COLORS["muted"], after=0)

    doc.core_properties.title = "Creatorshop Brand Guidelines Book"
    doc.core_properties.subject = "Brand guidelines based on the Creatorshop codebase"
    doc.core_properties.author = "OpenAI Codex"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
